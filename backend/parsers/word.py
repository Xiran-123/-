"""
Word文档解析器
支持 .docx 格式
"""
from typing import List, Dict
import os


def parse_word(file_path: str) -> List[Dict]:
    """解析Word文档，返回段落列表"""
    ext = file_path.rsplit(".", 1)[-1].lower()
    
    if ext == "docx":
        return _parse_docx(file_path)
    elif ext == "doc":
        return _parse_doc(file_path)
    else:
        return []


def _parse_docx(file_path: str) -> List[Dict]:
    """解析.docx文件"""
    try:
        from docx import Document
    except ImportError:
        raise RuntimeError("请安装python-docx: pip install python-docx")
    
    doc = Document(file_path)
    paragraphs = []
    
    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if text:
            paragraphs.append({
                "index": i,
                "style": para.style.name if para.style else "",
                "content": text,
            })
    
    # 提取表格内容
    for table_idx, table in enumerate(doc.tables):
        for row_idx, row in enumerate(table.rows):
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                paragraphs.append({
                    "index": len(paragraphs),
                    "style": f"Table_{table_idx}_Row_{row_idx}",
                    "content": " | ".join(cells),
                })
    
    return paragraphs


def _parse_doc(file_path: str) -> List[Dict]:
    """解析.doc文件（需要antiword或libreoffice，这里做降级处理）"""
    # 尝试用docx兼容（部分.doc实际上是docx）
    try:
        return _parse_docx(file_path)
    except Exception:
        pass
    
    # 尝试文本提取
    paragraphs = []
    try:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            text = f.read()
        # 简单过滤
        lines = [l.strip() for l in text.split("\n") if l.strip()]
        for i, line in enumerate(lines):
            paragraphs.append({
                "index": i,
                "style": "Text",
                "content": line,
            })
    except Exception as e:
        paragraphs.append({
            "index": 0,
            "style": "Error",
            "content": f"无法解析.doc文件: {str(e)}。建议另存为.docx格式。",
        })
    
    return paragraphs
